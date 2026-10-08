from docx import Document
d=Document(); d.add_heading('ТЗ: Корзина',1); d.add_paragraph('2.1. Пользователь добавляет товар в корзину кнопкой «В корзину».')
t=d.add_table(rows=2,cols=2); t.cell(0,0).text='Поле'; t.cell(0,1).text='Ограничение'; t.cell(1,0).text='Количество'; t.cell(1,1).text='от 1 до 99'
d.save('t/cart.docx')
import openpyxl
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Требования'; ws.append(['ID','Требование']); ws.append(['FR-1','Сумма заказа пересчитывается при изменении количества']); wb.save('t/req.xlsx')
from reportlab.pdfgen import canvas; from reportlab.pdfbase import pdfmetrics; from reportlab.pdfbase.ttfonts import TTFont
pdfmetrics.registerFont(TTFont('DV','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
c=canvas.Canvas('t/spec.pdf'); c.setFont('DV',12); c.drawString(72,760,'3.1. POST /api/v1/cart возвращает 201 Created'); c.drawString(72,740,'Ошибка 409 при превышении лимита'); c.save()
