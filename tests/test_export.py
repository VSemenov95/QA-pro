"""Выгрузка результата: стандартный JSON, файл по описанию формата (CSV, XLSX, JSON, Markdown, XML), JSON для Сферы ФТ; скачивание в артефакте и в автономной версии."""
import base64
import io
import json
import re

import pytest

from tests.support import files
from tests.support.model_generation import AI_CONN, connect_ai, standalone, three_steps_cases

CSV_SPEC = {'fileType': 'csv', 'filename': 'testit-import', 'mode': 'map', 'rowMode': 'step', 'delimiter': ';', 'notes': 'CSV для Test IT, шаг на строку.',
            'cols': [{'header': 'ID', 'value': '{{id}}', 'firstOnly': True}, {'header': 'Название', 'value': '{{title}}', 'firstOnly': True},
                     {'header': 'Предусловия', 'value': '{{preconditions|; }}', 'firstOnly': True}, {'header': 'Шаг', 'value': '{{step.n}}. {{step.action}}'},
                     {'header': 'Ожидаемый результат', 'value': '{{step.expected}}'}, {'header': 'Приоритет', 'value': '{{priority=high:Высокий,medium:Средний,low:Низкий}}', 'firstOnly': True},
                     {'header': 'Теги', 'value': '{{suites|, }}', 'firstOnly': True}], 'itemRules': 'по шаблону'}
FORMAT = 'CSV для импорта в Test IT, шаг на строку'


@pytest.fixture
def done(app):
    app.add_sample()
    app.generate()
    return app


@pytest.fixture
def steps3(app):
    app.model.on('promptGenerate', three_steps_cases)
    app.add_sample()
    app.generate()
    return app


def describe_format(app, text=FORMAT):
    app.open_settings('export')
    app.page.fill('#cfg-export', text)
    app.close_settings()
    app.tab('export')


def build(app, wait='[data-act=dl-out]'):
    app.page.click('[data-act=run-export]')
    app.page.wait_for_selector(wait)


def download(app):
    n = len(app.downloads())
    app.page.click('[data-act=dl-out]')
    app.page.wait_for_function('(n) => (window.__downloads || []).length > n', arg=n)
    return app.downloads()[-1]


def blob_bytes(app):
    b64 = app.js("""async () => { const b = window.__dlBlob; const u = new Uint8Array(await b.arrayBuffer()); let s = ''; for (const x of u) s += String.fromCharCode(x); return btoa(s); }""")
    return base64.b64decode(b64)


def created_day(app):
    return app.state('new Date(S.run.createdAt).toISOString().slice(0, 10)')


# ---------- стандартный JSON ----------
def test_export_tab_is_empty_before_generation(app):
    app.tab('export')
    assert 'Экспорт появится после генерации тест-кейсов' in app.text('#tab-body')


def test_default_export_is_offered_without_format(done):
    done.tab('export')
    assert 'Стандартный JSON' in done.text('.exp-card') and done.count('[data-act=run-export]') == 0
    assert done.text('.sec-h h3') == f'test-cases_{created_day(done)}.json'


def test_default_export_file(done):
    done.tab('export')
    d = download(done)
    assert d['filename'] == f'test-cases_{created_day(done)}.json'
    out = json.loads(d['text'])
    assert out['sources'] == ['Пример — ТЗ восстановление пароля.md']
    assert {k: [c['id'] for c in v] for k, v in out['testCases'].items()} == {'api': ['API-001', 'API-002'], 'ui': ['UI-001'], 'e2e': ['E2E-001']}


def test_default_export_suites_hold_case_numbers(done):
    done.tab('export')
    out = json.loads(download(done)['text'])
    assert out['suites'] == {'smoke': ['API-001', 'UI-001', 'E2E-001'], 'regression': ['API-001', 'API-002', 'UI-001', 'E2E-001']}


def test_default_export_traceability(done):
    done.tab('export')
    out = json.loads(download(done)['text'])
    assert [(t['requirement'], t['testCases']) for t in out['traceability']] == [('R-001', ['UI-001']), ('R-002', ['API-001', 'E2E-001']), ('R-003', ['E2E-001']), ('R-004', ['API-002'])]
    assert out['traceability'][0]['title'] == 'Ссылка «Забыли пароль?»'


def test_default_export_has_no_actualization_block_for_generation(done):
    done.tab('export')
    assert 'actualization' not in json.loads(download(done)['text'])


def test_export_of_actualization_has_statuses(app, tmp_path):
    app.mode('update')
    app.upload_tests(files.cases_json(tmp_path))
    app.upload(files.spec_file(tmp_path, version='2.0'))
    app.generate()
    app.tab('export')
    a = json.loads(download(app)['text'])['actualization']
    assert a['source'] == 'cases.json' and a['summary'] == {'same': 6, 'changed': 1, 'new': a['summary']['new'], 'obsolete': 1}
    by_id = {c['id']: c for c in a['cases']}
    assert by_id['UI-001']['status'] == 'changed' and by_id['UI-001']['statusRu'] == 'Изменён' and 'Введите корректный email' in by_id['UI-001']['changes'][0]
    assert by_id['UI-002']['status'] == 'obsolete' and by_id['UI-002']['sourceId'] == 'UI-002' and 'SMS' in by_id['UI-002']['reason']


def test_describe_format_button_opens_export_settings(done):
    done.tab('export')
    done.page.click('[data-act=cfg-export]')
    assert done.visible('#set-ov') and done.page.is_visible('#cfg-export')


# ---------- файл по описанию формата ----------
def test_format_description_goes_to_the_model_once(done):
    describe_format(done)
    build(done)
    prompts = done.model.prompts_of('promptExportSpec')
    assert len(prompts) == 1 and FORMAT in prompts[0] and 'API-001' in prompts[0]


def test_csv_export_when_fields_can_be_mapped_has_no_other_model_calls(done):
    done.model.on('promptExportSpec', CSV_SPEC)
    describe_format(done)
    build(done)
    assert len(done.model.calls_of('promptExportSpec')) == 1 and len(done.model.calls_of('promptExportItems')) == 0
    assert 'Файл собирается по шаблону полей, без модели.' in done.text('.exp-card')


def test_csv_file_content(done):
    done.model.on('promptExportSpec', CSV_SPEC)
    describe_format(done)
    build(done)
    d = download(done)
    assert d['filename'] == 'testit-import.csv'
    lines = d['text'].split('\r\n')
    assert lines[0] == '﻿ID;Название;Предусловия;Шаг;Ожидаемый результат;Приоритет;Теги'
    assert lines[1] == 'API-001;[BE] Кейс для P-02: проверка восстановления пароля;Пользователь зарегистрирован;1. Открыть /login;Открыта страница;Высокий;smoke, regression'
    assert len(lines) == 5 and lines[2].startswith('API-002;') and lines[2].endswith(';regression')
    assert 'ID;Название;Предусловия;Шаг;Ожидаемый результат;Приоритет;Теги' in done.text('.json-pre')   # предпросмотр показывает тот же файл


def test_csv_step_per_row_leaves_case_cells_empty_after_first_row(steps3):
    steps3.model.on('promptExportSpec', CSV_SPEC)
    describe_format(steps3)
    build(steps3)
    lines = download(steps3)['text'].split('\r\n')[1:5]
    assert lines[0] == 'API-001;[BE] Три шаг;Пользователь зарегистрирован;1. Действие 1;Результат 1;Высокий;smoke, regression'.replace('Три шаг', 'Три шага')
    assert lines[1:3] == [';;;2. Действие 2;Результат 2;;', ';;;3. Действие 3;Результат 3;;']
    assert lines[3].startswith('API-002;')


def test_csv_cell_with_delimiter_and_quotes_is_quoted(done):
    done.model.on('promptExportSpec', {**CSV_SPEC, 'cols': [{'header': 'ID', 'value': '{{id}}'}, {'header': 'Название', 'value': '{{title}}'}], 'rowMode': 'case'})
    done.tab('cases')
    done.page.click('#case-list article:has(.case-id:text-is("API-001")) .case-h')
    done.page.click('#case-list [data-act=form-case]')
    done.page.fill('.case-form input[data-fp="title"]', '[BE] Название; с "кавычками"')
    done.page.click('[data-act=ff-save]')
    describe_format(done)
    build(done)
    assert download(done)['text'].split('\r\n')[1] == 'API-001;"[BE] Название; с ""кавычками"""'


def test_changed_case_makes_the_file_stale_and_rebuild_needs_no_new_format(done):
    done.model.on('promptExportSpec', CSV_SPEC)
    describe_format(done)
    build(done)
    done.tab('cases')
    done.page.click('#case-list article:has(.case-id:text-is("API-001")) .case-h')
    done.page.click('#case-list [data-act=form-case]')
    done.page.fill('.case-form input[data-fp="title"]', '[BE] Новое название')
    done.page.click('[data-act=ff-save]')
    done.tab('export')
    assert 'Кейсы изменились после сборки — пересоберите файл.' in done.text('.banner.warn')
    done.page.click('[data-act=run-export]')
    done.page.wait_for_function("() => !document.querySelector('.banner.warn')")
    assert 'API-001;[BE] Новое название;' in download(done)['text']
    assert len(done.model.calls_of('promptExportSpec')) == 1


def test_rebuild_from_scratch_asks_for_format_again(done):
    done.model.on('promptExportSpec', CSV_SPEC)
    describe_format(done)
    build(done)
    done.page.click('[data-act=rebuild-export]')
    done.page.wait_for_function("() => window.__qa.S.exp && window.__qa.S.exp.status === 'done'")
    assert len(done.model.calls_of('promptExportSpec')) == 2


def test_other_format_description_builds_another_file(done):
    done.model.on('promptExportSpec', CSV_SPEC)
    describe_format(done)
    build(done)
    done.open_settings('export')
    done.page.fill('#cfg-export', 'То же, но через запятую')
    done.close_settings()
    done.model.on('promptExportSpec', {**CSV_SPEC, 'filename': 'comma', 'delimiter': ','})
    done.tab('export')
    done.page.click('[data-act=run-export]')
    done.page.wait_for_function("() => window.__qa.S.exp.spec && window.__qa.S.exp.spec.filename === 'comma'")
    assert download(done)['filename'] == 'comma.csv'


AI_SPEC = {'fileType': 'csv', 'filename': 'ai', 'mode': 'ai', 'columns': ['ID', 'Название'], 'delimiter': ',', 'itemRules': 'ID и название'}


def case_ids(prompt):
    return re.findall(r'"id": "([A-Z0-9-]+)",\n\s+"type"', prompt)


def rows_for(call, skip=()):
    ids = case_ids(call.prompt)
    return {'items': [{'caseId': i, 'rows': [[i, 'Название ' + i]]} for i in ids if i not in skip]}


def test_ai_mode_asks_model_per_batch_of_cases(done):
    done.model.on('promptExportSpec', AI_SPEC).on('promptExportItems', rows_for)
    describe_format(done)
    build(done)
    assert len(done.model.calls_of('promptExportItems')) == 1
    assert download(done)['text'].split('\r\n') == ['\ufeffID,Название', 'API-001,Название API-001', 'API-002,Название API-002', 'UI-001,Название UI-001', 'E2E-001,Название E2E-001']


def test_ai_mode_batches_of_eight(app, tmp_path):
    from tests.support import seed
    cases = [seed.case(f'u{i}', 'api', f'API-{i:03d}', f'[BE] POST /x — проверка {i}', ['R-001']) for i in range(1, 20)]
    run = seed.run('run_x', 'Док', cases, [seed.req('R-001', 'Док')], stage='draft')
    app.import_product(seed.write(tmp_path, seed.product_export(runs=[run])))
    app.page.wait_for_selector('#pcase-list .sec', timeout=15000)
    app.call("a.openHistory('run_x', { tab: 'export' })")
    app.model.on('promptExportSpec', AI_SPEC).on('promptExportItems', rows_for)
    describe_format(app)
    build(app)
    assert sorted(len(case_ids(c.prompt)) for c in app.model.calls_of('promptExportItems')) == [3, 8, 8]


def test_cases_the_model_skipped_are_reported_and_repeated_alone(done):
    done.model.on('promptExportSpec', AI_SPEC).on('promptExportItems', lambda c: rows_for(c, skip=('API-002',)))
    describe_format(done)
    build(done)
    assert 'Не удалось преобразовать: API-002.' in done.text('.banner.warn')
    assert 'API-002' not in download(done)['text']
    done.model.on('promptExportItems', rows_for)
    done.page.click('[data-act=run-export]')
    done.page.wait_for_function("() => window.__qa.S.exp.status === 'done' && !window.__qa.S.exp.error")
    assert case_ids(done.model.calls_of('promptExportItems')[-1].prompt) == ['API-002']
    assert 'API-002,Название API-002' in download(done)['text']


def test_csv_without_columns_is_an_error(done):
    done.model.on('promptExportSpec', {'fileType': 'csv', 'mode': 'ai', 'filename': 'x'})
    describe_format(done)
    done.page.click('[data-act=run-export]')
    done.page.wait_for_selector('.banner.err')
    assert 'Модель не определила колонки файла' in done.text('.banner.err')


def test_json_by_description_uses_sets_and_item_template(done):
    done.model.on('promptExportSpec', {'fileType': 'json', 'filename': 'cases', 'mode': 'map',
                                       'wrapper': {'project': 'Пример', 'created': '%%GENERATED_AT%%', 'smoke': '%%SMOKE%%', 'count': 'x'},
                                       'item': {'id': '{{id}}', 'name': '{{title}}', 'steps': {'$each': 'steps', 'item': '{{action}} → {{expected}}'}}})
    describe_format(done, 'JSON: проект и набор смоук')
    build(done)
    d = download(done)
    out = json.loads(d['text'])
    assert d['filename'] == 'cases.json' and out['project'] == 'Пример' and out['created'].startswith(created_day(done))
    assert [c['id'] for c in out['smoke']] == ['API-001', 'UI-001', 'E2E-001']
    assert out['smoke'][0] == {'id': 'API-001', 'name': '[BE] Кейс для P-02: проверка восстановления пароля', 'steps': ['Открыть /login → Открыта страница']}


def test_markdown_by_description(done):
    done.model.on('promptExportSpec', {'fileType': 'md', 'filename': 'cases', 'mode': 'map', 'wrapper': '# Кейсы\n\n%%ALL%%', 'separator': '\n\n', 'itemText': '## {{id}} {{title}}\n{{#steps}}- {{action}}\n{{/steps}}'})
    describe_format(done, 'Markdown со списком шагов')
    build(done)
    d = download(done)
    assert d['filename'] == 'cases.md' and d['text'].startswith('# Кейсы\n\n## API-001 [BE] Кейс для P-02')
    assert '- Открыть /login\n' in d['text'] and d['text'].count('## ') == 4


def test_xml_is_saved_as_txt(done):
    done.model.on('promptExportSpec', {'fileType': 'xml', 'filename': 'cases', 'mode': 'map', 'wrapper': '<cases>\n%%ALL%%\n</cases>', 'itemText': '<case id="{{id}}"/>'})
    describe_format(done, 'XML')
    build(done)
    d = download(done)
    assert d['filename'] == 'cases.xml.txt' and d['text'] == '<cases>\n<case id="API-001"/>\n<case id="API-002"/>\n<case id="UI-001"/>\n<case id="E2E-001"/>\n</cases>'


@pytest.mark.needs_libs
def test_xlsx_by_description(done):
    import openpyxl
    done.model.on('promptExportSpec', {**CSV_SPEC, 'fileType': 'xlsx', 'filename': 'cases-xlsx', 'rowMode': 'case', 'cols': CSV_SPEC['cols'][:2]})
    describe_format(done, 'Excel для Test IT')
    build(done)
    assert done.count('.steps-t th') == 2   # предпросмотр таблицей
    d = download(done)
    assert d['filename'] == 'cases-xlsx.xlsx'
    ws = openpyxl.load_workbook(io.BytesIO(blob_bytes(done))).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ['ID', 'Название'] and rows[1] == ['API-001', '[BE] Кейс для P-02: проверка восстановления пароля'] and len(rows) == 5


@pytest.mark.needs_libs
def test_xlsx_cells_have_code_without_backticks(app):
    """В XLSX блок кода остаётся текстом без тройных кавычек."""
    import openpyxl
    def gen(call):
        res = three_steps_cases(call)
        for c in res['cases']:
            c['testCase']['steps'][0]['expected'] = 'Тело ответа:\n```json\n{"status":"sent"}\n```'
        return res
    app.model.on('promptGenerate', gen).on('promptExportSpec', {**CSV_SPEC, 'fileType': 'xlsx', 'filename': 'x', 'rowMode': 'step', 'cols': [{'header': 'ID', 'value': '{{id}}', 'firstOnly': True}, {'header': 'Результат', 'value': '{{step.expected}}'}]})
    app.add_sample()
    app.generate()
    describe_format(app, 'Excel')
    build(app)
    download(app)
    ws = openpyxl.load_workbook(io.BytesIO(blob_bytes(app))).active
    assert ws['B2'].value == 'Тело ответа:\n{"status":"sent"}'


# ---------- JSON для Сферы ФТ ----------
def sfera_download(app):
    app.tab('cases')
    n = len(app.downloads())
    app.page.click('[data-act=sfera-json]')
    app.page.wait_for_function('(n) => (window.__downloads || []).length > n', arg=n)
    d = app.downloads()[-1]
    return d['filename'], json.loads(d['text'])


def test_sfera_json_file(done):
    name, out = sfera_download(done)
    assert name == 'qa-pro-sfera.json' and out['format'] == 'qa-pro/sfera-v1' and out['product'] == 'Мой продукт'
    assert [c['qaId'] for c in out['cases']] == ['API-001', 'API-002', 'UI-001', 'E2E-001']


def test_sfera_case_fields(done):
    _, out = sfera_download(done)
    c = out['cases'][0]
    assert set(c) == {'ref', 'qaId', 'name', 'description', 'preConditions', 'steps', 'postConditions', 'priority', 'suites'}
    assert c['name'] == '[BE] Кейс для P-02: проверка восстановления пароля' and c['priority'] == 'P3_HIGH' and c['suites'] == ['smoke', 'regression']
    assert c['preConditions'] == [{'action': '<p>Пользователь зарегистрирован</p>', 'expectedResult': '<p></p>'}]
    assert c['steps'] == [{'action': '<p>Открыть /login</p>', 'expectedResult': '<p>Открыта страница</p>'}]
    assert c['ref'] == done.state("S.run.cases[0].uid")


def test_sfera_json_button_counts_cases(done):
    done.tab('cases')
    assert done.text('[data-act=sfera-json]') == 'Скачать JSON для загрузчика (4)'


@pytest.mark.parametrize('prio,expected', [('High', 'P3_HIGH'), ('Medium', 'P2_MEDIUM'), ('Low', 'P1_LOW'), ('Высокий', 'P3_HIGH'), ('Низкий', 'P1_LOW'), ('', 'P2_MEDIUM'), ('Blocker', 'P2_MEDIUM')])
def test_sfera_priority(app, prio, expected):
    assert app.call("a.sferaCase({ uid: 'u', dispId: 'API-001', type: 'api', suites: [], requirementIds: [], testCase: { title: 'x', priority: arg, steps: [] } }).priority", prio) == expected


def sfera(app, tc):
    return app.call("a.sferaCase({ uid: 'u1', dispId: 'API-001', type: 'api', suites: ['smoke'], requirementIds: [], testCase: arg })", tc)


def test_sfera_code_block_becomes_pre_with_empty_paragraphs(app):
    """README: пояснение → <p>, пустой <p>, <pre data-id style>, пустой <p>."""
    c = sfera(app, {'title': 'Кейс', 'priority': 'High', 'steps': [{'action': 'Отправить запрос', 'expected': 'Тело ответа:\n```json\n{"status":"sent"}\n```'}]})
    exp = c['steps'][0]['expectedResult']
    assert re.fullmatch(r'<p>Тело ответа:</p>\n<p></p>\n<pre data-id="[0-9a-f-]{36}" style="height: 48px">\n\{"status":"sent"\}</pre>\n<p></p>', exp), exp


def test_sfera_pre_height_follows_line_count_and_escapes_markup(app):
    c = sfera(app, {'title': 'Кейс', 'steps': [{'action': '```xml\n<a>1</a>\n<b>&</b>\n```', 'expected': ''}]})
    a = c['steps'][0]['action']
    assert 'style="height: 72px">\n&lt;a&gt;1&lt;/a&gt;\n&lt;b&gt;&amp;&lt;/b&gt;</pre>' in a


def test_sfera_json_inside_text_is_also_code(app):
    c = sfera(app, {'title': 'Кейс', 'steps': [{'action': 'Тело запроса: {"email": "user@example.ru"}', 'expected': ''}]})
    assert '<pre data-id=' in c['steps'][0]['action'] and '{"email": "user@example.ru"}</pre>' in c['steps'][0]['action']


def test_sfera_plain_text_is_escaped_and_lines_are_broken(app):
    c = sfera(app, {'title': 'Кейс', 'steps': [{'action': 'Ввести <script> & нажать\nОтправить', 'expected': ''}]})
    assert c['steps'][0]['action'] == '<p>Ввести &lt;script&gt; &amp; нажать<br>Отправить</p>' and c['steps'][0]['expectedResult'] == '<p></p>'


def test_sfera_name_is_cut_to_256(app):
    assert len(sfera(app, {'title': 'я' * 300})['name']) == 256


def test_sfera_json_skips_obsolete_cases(app, tmp_path):
    app.mode('update')
    app.upload_tests(files.cases_json(tmp_path))
    app.upload(files.spec_file(tmp_path, version='2.0'))
    app.generate()
    app.tab('cases')
    total = app.state('S.run.cases.length')
    assert app.text('[data-act=sfera-json]') == f'Скачать JSON для загрузчика ({total - 1})'
    assert 'UI-002' not in [c['qaId'] for c in sfera_download(app)[1]['cases']]


# ---------- автономная версия: файлы скачиваются якорем ----------
@pytest.fixture
def alone(make_app):
    app = make_app(**standalone())
    connect_ai(app)
    app.add_sample()
    app.generate()
    return app


def anchor_download(app, selector):
    with app.page.expect_download() as dl:
        app.page.click(selector)
    d = dl.value
    with open(d.path(), 'rb') as f:
        return d.suggested_filename, f.read()


def test_standalone_generation_goes_through_the_connection(alone):
    assert alone.state('S.run.status') == 'done' and alone.state('S.run.cases.length') == 4
    assert alone.state("S.settings.genAI") == 'conn:' + AI_CONN['id']


def test_standalone_default_export_download(alone):
    alone.tab('export')
    name, data = anchor_download(alone, '[data-act=dl-out]')
    assert name == f'test-cases_{created_day(alone)}.json'
    assert [c['id'] for c in json.loads(data.decode('utf-8'))['testCases']['api']] == ['API-001', 'API-002']


def test_standalone_checklist_download(alone):
    alone.tab('checklist')
    name, data = anchor_download(alone, '[data-act=cl-dl-md]')
    assert name == f'checklist_{created_day(alone)}.md' and data.decode('utf-8').startswith('# Чек-лист\n\n## Форма восстановления')


def test_standalone_sfera_json_download(alone):
    alone.tab('cases')
    name, data = anchor_download(alone, '[data-act=sfera-json]')
    out = json.loads(data.decode('utf-8'))
    assert name == 'qa-pro-sfera.json' and out['format'] == 'qa-pro/sfera-v1' and len(out['cases']) == 4


def test_standalone_csv_download_has_bom_for_excel(alone):
    alone.model.on('promptExportSpec', CSV_SPEC)
    describe_format(alone)
    build(alone)
    name, data = anchor_download(alone, '[data-act=dl-out]')
    assert name == 'testit-import.csv' and data.startswith(b'\xef\xbb\xbfID;')


@pytest.mark.needs_libs
def test_standalone_xlsx_download(alone):
    import openpyxl
    alone.model.on('promptExportSpec', {**CSV_SPEC, 'fileType': 'xlsx', 'filename': 'cases', 'rowMode': 'case', 'cols': CSV_SPEC['cols'][:2]})
    describe_format(alone, 'Excel')
    build(alone)
    name, data = anchor_download(alone, '[data-act=dl-out]')
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert name == 'cases.xlsx' and [c.value for c in ws[1]] == ['ID', 'Название'] and ws.max_row == 5


def test_standalone_export_description_survives_reload(alone):
    describe_format(alone)
    alone.page.wait_for_function("() => (localStorage.getItem('qapro.cfgDraft') || '').includes('Test IT')")   # черновик указаний пишется с задержкой
    alone.reload()
    assert alone.state('S.cfg.exportFormat') == FORMAT and alone.state('S.cfgDirty.export') is True
