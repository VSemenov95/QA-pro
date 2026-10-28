"""Загруженный чек-лист: из его пунктов пишутся тест-кейсы, документы остаются контекстом."""
import json

import pytest


def checklist_file(tmp_path):
    items = [
        {'check': 'Ссылка «Забыли пароль?» на /login', 'expected': 'Ссылка видна под полем пароля', 'kind': 'positive', 'type': 'ui', 'priority': 'P3_HIGH', 'suites': ['smoke', 'regression'], 'requirementIds': []},
        {'check': 'Email без @ в форме восстановления', 'expected': 'Сообщение об ошибке, запрос не отправлен', 'kind': 'negative', 'type': 'ui', 'priority': 'P2_MEDIUM', 'suites': ['regression'], 'requirementIds': []},
    ]
    p = tmp_path / 'checklist.json'
    p.write_text(json.dumps({'groups': [{'title': 'Восстановление пароля', 'items': items}]}, ensure_ascii=False), encoding='utf-8')
    return p


def test_uploaded_checklist_becomes_the_basis_of_cases(app, tmp_path):
    app.add_sample()
    app.page.set_input_files('#cl-upload', str(checklist_file(tmp_path)))
    app.page.wait_for_function("() => /пункт/.test(document.querySelector('#cl-upload-info').textContent)", timeout=10000)
    app.generate()
    checks = app.state("S.run.checklist.groups.flatMap((g) => g.items.map((i) => i.check))")
    assert sorted(checks) == sorted(['Ссылка «Забыли пароль?» на /login', 'Email без @ в форме восстановления']), checks
    plan = ' '.join(app.state("S.run.plan.flatMap((p) => p.checks)"))
    assert 'Забыли пароль' in plan and 'Email без @' in plan, plan   # пункты чек-листа легли в план кейсов


def test_clearing_the_uploaded_checklist_restores_the_default(app, tmp_path):
    app.page.set_input_files('#cl-upload', str(checklist_file(tmp_path)))
    app.page.wait_for_function("() => /пункт/.test(document.querySelector('#cl-upload-info').textContent)", timeout=10000)
    app.page.click('#cl-upload-clear')
    assert app.state('S.uploadedChecklist') is None


CHECKS = ['Ссылка «Забыли пароль?» на /login', 'Email без @ в форме восстановления']
TABLE = [
    ['№', 'Название', 'Приоритет', 'Компонент', 'Ожидаемый результат'],
    ['1', CHECKS[0], 'высокий', 'FE', 'Ссылка видна под полем пароля'],
    ['2', CHECKS[1], 'средний', 'FE', 'Сообщение об ошибке, запрос не отправлен'],
]


def assert_loaded(app):
    app.page.wait_for_function("() => /Чек-лист «/.test(document.querySelector('#cl-upload-info').textContent)", timeout=10000)
    assert sorted(app.state("S.uploadedChecklist.groups.flatMap((g) => g.items.map((i) => i.check))")) == sorted(CHECKS)


def test_markdown_table_from_export_is_a_checklist(app, tmp_path):
    p = tmp_path / 'checklist.md'
    p.write_text('# Чек-лист\n\n' + '\n'.join('| ' + ' | '.join(r) + ' |' if i else '| ' + ' | '.join(r) + ' |\n|---:|---|---|---|---|' for i, r in enumerate(TABLE)) + '\n', encoding='utf-8')
    app.page.set_input_files('#cl-upload', str(p))
    assert_loaded(app)
    assert app.state("S.uploadedChecklist.groups[0].items[0].priority") == 'P3_HIGH'


@pytest.mark.needs_libs
def test_docx_table_from_export_is_a_checklist(app, tmp_path):
    from docx import Document
    d = Document()
    t = d.add_table(rows=len(TABLE), cols=5)
    for i, row in enumerate(TABLE):
        for j, cell in enumerate(row):
            t.cell(i, j).text = cell
    p = tmp_path / 'checklist.docx'
    d.save(str(p))
    app.page.set_input_files('#cl-upload', str(p))
    assert_loaded(app)
